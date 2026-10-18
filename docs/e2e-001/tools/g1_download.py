"""NEKKORA-E2E-001 · G1 — leitura do desafio e download REAL do Excel com o nekkora_rpa (operador: TEST_AGENT).

Abre https://rpachallenge.com num Chrome real isolado (contexto novo, sem perfil), no ambiente ``benchmark`` (origem
exata; recursos de terceiros bloqueados), clica no link de download pela CHAVE do registro, valida o arquivo
(origem, nome de destino, tamanho, assinatura XLSX), grava de forma atômica e lê com ``nekkora_rpa.tabular``.

Saídas:
- ``--evidencia``: JSON SEM VALORES (hash, tamanho, cabeçalhos, contagem, tipos por coluna, bloqueios de rede);
- ``--transiente``: JSON com o 1º registro (valores demonstrativos), usado SÓ para digitar na gravação do Studio;
  fica fora dos repositórios e é apagado pela homologação ao terminar.
Não clica em Start.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from nekkora_rpa import (
    Actions,
    BrowserSession,
    EventLogger,
    ExecutionContext,
    MemoryEventSink,
    SelectorRegistry,
)
from nekkora_rpa.core.config import from_mapping
from nekkora_rpa.tabular import Column, TabularSchema, read_xlsx

ORIGEM = "https://rpachallenge.com"
REGISTRO = {
    "schema": "nekkora-rpa/selectors", "version": "1.1.0", "system": "g1",
    "targets": {"g1.inicio.baixar": {"preferred": {"strategy": "css", "expression": "a[href$=\"downloadFiles/challenge.xlsx\"]"},
                                     "expected": {"tag": "a"}}},
}
ROTULOS = ["First Name", "Last Name", "Company Name", "Role in Company", "Address", "Email", "Phone Number"]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--destino", required=True, type=Path)
    ap.add_argument("--evidencia", required=True, type=Path)
    ap.add_argument("--transiente", required=True, type=Path)
    ap.add_argument("--headless", action="store_true")
    a = ap.parse_args()
    config = from_mapping({
        "runtime": {"environment": "benchmark", "allowed_origins": [ORIGEM], "base_url": ORIGEM + "/", "process_name": "g1_download"},
        "timeouts": {"action_ms": 15000, "navigation_ms": 30000, "transaction_ms": 60000},
        "browser": {"channel": "chrome", "headless": a.headless, "accept_downloads": True},
    }, base_dir=Path.cwd())
    eventos = MemoryEventSink()
    reg = SelectorRegistry.from_text(json.dumps(REGISTRO), environment="benchmark")
    ctx = ExecutionContext.create(config, EventLogger("ex-g1", "0.2.0", [eventos], reg.revision))
    sessao = BrowserSession(config.browser, config.origins, config.timeouts)
    sessao.start()
    try:
        act = Actions(ctx, sessao, reg)
        act.navigate("/", step_id="G1-ABRIR")
        instrucoes = act.page.inner_text("body")
        a.destino.mkdir(parents=True, exist_ok=True)
        f = act.download("g1.inicio.baixar", dest_dir=a.destino, file_name="challenge.xlsx", magic=b"PK\x03\x04", step_id="G1-BAIXAR")
        bloqueios = {"requisicoes_bloqueadas": sessao.blocked_requests, "navegacoes_bloqueadas": sessao.blocked_navigations}
    finally:
        sessao.close()
    esquema = TabularSchema(tuple(Column(r, integer_as_text=(r == "Phone Number")) for r in ROTULOS))
    dados = read_xlsx(f.path, esquema)
    import openpyxl  # tipos ORIGINAIS das células (sem valores) para a evidência

    wb = openpyxl.load_workbook(f.path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    linhas = list(ws.iter_rows(values_only=True))
    tipos: dict[str, list[str]] = {}
    for i, h in enumerate(linhas[0]):
        tipos[str(h)] = sorted({type(r[i]).__name__ for r in linhas[1:] if r and any(v is not None for v in r)})
    abas = wb.sheetnames
    wb.close()
    evid = {
        "origem": ORIGEM,
        "instrucoes_confirmadas": all(t in instrucoes for t in ("10 rounds", "Start", "spreadsheet")),
        "arquivo": {"nome_destino": f.path.name, "bytes": f.size_bytes, "sha256": f.sha256},
        "abas": abas,
        "cabecalhos_brutos": [repr(h) for h in linhas[0]],
        "cabecalhos_normalizados": list(dados.headers),
        "registros": len(dados.rows),
        "tipos_por_coluna": tipos,
        "linhas_com_problema": sum(1 for r in dados.rows if r.problems),
        "rede": bloqueios,
        "eventos": [e.event_code.value for e in eventos.events],
    }
    a.evidencia.parent.mkdir(parents=True, exist_ok=True)
    a.evidencia.write_text(json.dumps(evid, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    a.transiente.write_text(json.dumps({"registro1": dict(dados.rows[0].values)}, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({k: evid[k] for k in ("arquivo", "registros", "cabecalhos_normalizados", "linhas_com_problema", "rede")}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
